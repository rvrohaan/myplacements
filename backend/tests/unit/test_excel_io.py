"""app/services/excel_io.py - importing and exporting .xlsx.

This is where a college's real data enters the system, typed by hand into a
spreadsheet, so almost every input is wrong in some way: a CGPA in a text cell,
a "Yes" where a boolean belongs, data on the second tab, an empty row in the
middle. The parser's job is to keep going and report per-row, because one bad
row aborting a 400-student import is the failure mode that matters.

Workbooks are built in memory with openpyxl rather than read from fixtures, so
each test says exactly what shape it is about.
"""

import io

import pytest
from openpyxl import Workbook

from app.services.excel_io import Column, build_workbook, parse_rows

COLUMNS = [
    Column("Roll Number", "roll_number", required=True),
    Column("Name", "full_name", required=True),
    Column("Branch", "branch"),
    Column("CGPA", "cgpa", kind="float"),
    Column("Backlogs", "backlogs", kind="int"),
    Column("Higher Studies", "higher_studies_plan", kind="bool"),
    Column("Status", "placement_status", kind="enum", choices=["unplaced", "placed"]),
    Column("Computed", "computed", importable=False),
]


def workbook(*sheets: tuple[str, list[list]]) -> bytes:
    """A workbook from (title, rows) pairs. The first sheet is active."""
    wb = Workbook()
    wb.remove(wb.active)
    for title, rows in sheets:
        ws = wb.create_sheet(title=title)
        for row in rows:
            ws.append(row)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


HEADERS = ["Roll Number", "Name", "Branch", "CGPA", "Backlogs", "Higher Studies", "Status"]


def one_sheet(*rows: list) -> bytes:
    return workbook(("Sheet1", [HEADERS, *rows]))


# --- the happy path ---------------------------------------------------------


def test_a_clean_row_parses():
    parsed = parse_rows(one_sheet(["1RV001", "Asha", "CSE", 8.5, 0, "yes", "placed"]), COLUMNS)
    assert len(parsed) == 1
    assert parsed[0]["errors"] == []
    assert parsed[0]["data"] == {
        "roll_number": "1RV001",
        "full_name": "Asha",
        "branch": "CSE",
        "cgpa": 8.5,
        "backlogs": 0,
        "higher_studies_plan": True,
        "placement_status": "placed",
    }


def test_the_row_number_is_the_spreadsheet_row():
    """Reported back to the user, who is looking at the file. Off-by-one here
    sends them to the wrong line."""
    parsed = parse_rows(one_sheet(["1RV001", "Asha"], ["1RV002", "Bala"]), COLUMNS)
    assert [p["row"] for p in parsed] == [2, 3]


def test_headers_are_matched_case_insensitively():
    raw = workbook(("Sheet1", [["ROLL NUMBER", "name", "  Branch  "], ["1RV001", "Asha", "CSE"]]))
    assert parse_rows(raw, COLUMNS)[0]["data"]["full_name"] == "Asha"


def test_columns_may_appear_in_any_order():
    raw = workbook(("Sheet1", [["Name", "Roll Number"], ["Asha", "1RV001"]]))
    assert parse_rows(raw, COLUMNS)[0]["data"]["roll_number"] == "1RV001"


def test_unknown_columns_are_ignored():
    """People add their own working columns to the template all the time."""
    raw = workbook(("Sheet1", [["Roll Number", "Name", "Notes"], ["1RV001", "Asha", "call back"]]))
    parsed = parse_rows(raw, COLUMNS)
    assert parsed[0]["errors"] == []
    assert "Notes" not in parsed[0]["data"]


def test_an_export_only_column_is_not_imported():
    """importable=False means the sheet shows it but the parser must not accept
    it back - a computed field edited by hand would silently overwrite."""
    raw = workbook(("Sheet1", [["Roll Number", "Name", "Computed"], ["1RV001", "Asha", "junk"]]))
    assert "computed" not in parse_rows(raw, COLUMNS)[0]["data"]


# --- blank and partial rows -------------------------------------------------


def test_a_fully_blank_row_is_skipped():
    """Trailing blank rows are normal in a hand-edited sheet and must not
    become 300 spurious validation errors."""
    parsed = parse_rows(one_sheet(["1RV001", "Asha"], [None, None], ["1RV002", "Bala"]), COLUMNS)
    assert [p["data"]["roll_number"] for p in parsed] == ["1RV001", "1RV002"]


def test_a_row_of_empty_strings_is_skipped():
    parsed = parse_rows(one_sheet(["1RV001", "Asha"], ["", "   "]), COLUMNS)
    assert len(parsed) == 1


def test_a_blank_optional_cell_becomes_none():
    parsed = parse_rows(one_sheet(["1RV001", "Asha", None, None]), COLUMNS)
    assert parsed[0]["data"]["branch"] is None
    assert parsed[0]["data"]["cgpa"] is None
    assert parsed[0]["errors"] == []


def test_a_short_row_does_not_index_past_the_end():
    """openpyxl trims trailing empties, so the row can be shorter than the
    header. Reading past it would be an IndexError mid-import."""
    parsed = parse_rows(one_sheet(["1RV001", "Asha"]), COLUMNS)
    assert parsed[0]["errors"] == []
    assert parsed[0]["data"]["cgpa"] is None


# --- per-row validation -----------------------------------------------------


def test_a_bad_number_is_reported_against_its_row_not_raised():
    """The property that matters most: one bad row must not abort the batch."""
    parsed = parse_rows(one_sheet(["1RV001", "Asha", "CSE", "eight point five"]), COLUMNS)
    assert parsed[0]["errors"] == ["CGPA: expected a number, got 'eight point five'"]


def test_a_good_row_after_a_bad_one_still_parses():
    parsed = parse_rows(
        one_sheet(["1RV001", "Asha", "CSE", "junk"], ["1RV002", "Bala", "ECE", 9.1]), COLUMNS
    )
    assert parsed[0]["errors"]
    assert parsed[1]["errors"] == []
    assert parsed[1]["data"]["cgpa"] == 9.1


def test_a_missing_required_cell_is_reported():
    parsed = parse_rows(one_sheet([None, "Asha"]), COLUMNS)
    assert parsed[0]["errors"] == ["Roll Number is required"]


def test_the_error_names_the_header_the_user_sees():
    """Not the attribute name - they are looking at a spreadsheet, not at the
    model."""
    parsed = parse_rows(one_sheet(["1RV001", "Asha", "CSE", "junk"]), COLUMNS)
    assert parsed[0]["errors"][0].startswith("CGPA:")


def test_every_problem_in_a_row_is_collected():
    parsed = parse_rows(one_sheet([None, "Asha", "CSE", "junk", "junk"]), COLUMNS)
    assert len(parsed[0]["errors"]) == 3


# --- value coercion ---------------------------------------------------------


@pytest.mark.parametrize("raw", ["yes", "Yes", "YES", "y", "true", "1", True, 1])
def test_truthy_spellings_become_true(raw):
    """People write yes, Y and TRUE in the same column of the same file."""
    parsed = parse_rows(one_sheet(["1RV001", "Asha", "CSE", None, None, raw]), COLUMNS)
    assert parsed[0]["data"]["higher_studies_plan"] is True


@pytest.mark.parametrize("raw", ["no", "No", "n", "false", "0", False])
def test_falsy_spellings_become_false(raw):
    parsed = parse_rows(one_sheet(["1RV001", "Asha", "CSE", None, None, raw]), COLUMNS)
    assert parsed[0]["data"]["higher_studies_plan"] is False


def test_an_unrecognised_boolean_is_an_error():
    parsed = parse_rows(one_sheet(["1RV001", "Asha", "CSE", None, None, "maybe"]), COLUMNS)
    assert parsed[0]["errors"] == ["Higher Studies: expected yes/no, got 'maybe'"]


def test_a_whole_number_typed_as_a_decimal_is_accepted():
    """Excel stores 2 as 2.0, so int(float(...)) is doing real work here."""
    parsed = parse_rows(one_sheet(["1RV001", "Asha", "CSE", None, 2.0]), COLUMNS)
    assert parsed[0]["data"]["backlogs"] == 2


def test_a_number_in_a_text_cell_is_accepted():
    """A column formatted as text is the single most common import problem."""
    parsed = parse_rows(one_sheet(["1RV001", "Asha", "CSE", "8.5", "1"]), COLUMNS)
    assert parsed[0]["data"]["cgpa"] == 8.5
    assert parsed[0]["data"]["backlogs"] == 1


def test_strings_are_trimmed():
    parsed = parse_rows(one_sheet(["  1RV001  ", "  Asha  "]), COLUMNS)
    assert parsed[0]["data"]["roll_number"] == "1RV001"


def test_an_enum_is_lowercased():
    parsed = parse_rows(one_sheet(["1RV001", "Asha", "CSE", None, None, None, "PLACED"]), COLUMNS)
    assert parsed[0]["data"]["placement_status"] == "placed"


def test_an_enum_outside_its_choices_is_an_error_listing_them():
    parsed = parse_rows(one_sheet(["1RV001", "Asha", "CSE", None, None, None, "hired"]), COLUMNS)
    assert parsed[0]["errors"] == ["Status: 'hired' is not one of unplaced, placed"]


# --- choosing the sheet -----------------------------------------------------


def test_a_filled_second_tab_beats_an_empty_template_tab():
    """The regression the scoring exists for: taking whichever tab happened to
    be active silently imported nothing when the data sat on the second one."""
    raw = workbook(("Template", [HEADERS]), ("My Data", [HEADERS, ["1RV001", "Asha"]]))
    parsed = parse_rows(raw, COLUMNS)
    assert [p["data"]["roll_number"] for p in parsed] == ["1RV001"]


def test_the_sheet_naming_more_known_columns_wins():
    raw = workbook(
        ("Partial", [["Roll Number"], ["1RV999"]]),
        ("Full", [HEADERS, ["1RV001", "Asha", "CSE"]]),
    )
    assert parse_rows(raw, COLUMNS)[0]["data"]["full_name"] == "Asha"


def test_sheets_without_recognisable_headers_are_ignored():
    raw = workbook(("Notes", [["something", "else"], ["a", "b"]]), ("Data", [HEADERS, ["1RV001", "Asha"]]))
    assert len(parse_rows(raw, COLUMNS)) == 1


# --- file-level failures ----------------------------------------------------


def test_a_file_with_no_recognisable_headers_raises_with_guidance():
    """Raised, not returned empty: "0 imported" gives the user nothing to act
    on, and the message names the columns expected and the sheets checked."""
    raw = workbook(("Sheet1", [["Alpha", "Beta"], ["1", "2"]]))
    with pytest.raises(ValueError) as exc:
        parse_rows(raw, COLUMNS)
    assert "Roll Number" in str(exc.value)
    assert "Sheet1" in str(exc.value)


def test_a_missing_required_column_is_reported_once_for_the_file():
    """A header problem is one problem with the file, not the same problem
    repeated against 400 rows."""
    raw = workbook(("Sheet1", [["Roll Number", "Branch"], ["1RV001", "CSE"]]))
    with pytest.raises(ValueError) as exc:
        parse_rows(raw, COLUMNS)
    assert "Name" in str(exc.value)


def test_a_file_that_is_not_a_workbook_raises_a_readable_message():
    """Users upload .csv and .xls renamed to .xlsx."""
    with pytest.raises(ValueError, match="Could not read file as .xlsx"):
        parse_rows(b"this is not a spreadsheet", COLUMNS)


def test_an_empty_sheet_with_only_headers_parses_to_nothing():
    assert parse_rows(one_sheet(), COLUMNS) == []


# --- export -----------------------------------------------------------------


class Row:
    def __init__(self, **kw):
        self.__dict__.update(kw)


def test_the_export_writes_a_header_row_and_the_data():
    buffer = build_workbook(COLUMNS, [Row(roll_number="1RV001", full_name="Asha")], "Students")
    from openpyxl import load_workbook

    ws = load_workbook(buffer).active
    assert ws.title == "Students"
    assert [c.value for c in ws[1]][:2] == ["Roll Number", "Name"]
    assert [c.value for c in ws[2]][:2] == ["1RV001", "Asha"]


def test_the_export_unwraps_enum_members():
    """A raw PlacementStatus in a cell would render as the repr."""
    from app.models.student import PlacementStatus

    buffer = build_workbook(
        COLUMNS, [Row(roll_number="1RV001", placement_status=PlacementStatus.PLACED)], "Students"
    )
    from openpyxl import load_workbook

    ws = load_workbook(buffer).active
    assert ws.cell(row=2, column=7).value == "placed"


def test_a_missing_attribute_exports_as_blank_rather_than_raising():
    buffer = build_workbook(COLUMNS, [Row(roll_number="1RV001")], "Students")
    from openpyxl import load_workbook

    assert load_workbook(buffer).active.cell(row=2, column=2).value is None


def test_an_export_round_trips_back_through_the_parser():
    """The same Column list drives both, so a change that breaks the pairing
    should fail here rather than in a user's import next week."""
    rows = [Row(roll_number="1RV001", full_name="Asha", branch="CSE", cgpa=8.5, backlogs=0)]
    buffer = build_workbook(COLUMNS, rows, "Students")
    parsed = parse_rows(buffer.getvalue(), COLUMNS)
    assert parsed[0]["errors"] == []
    assert parsed[0]["data"]["roll_number"] == "1RV001"
    assert parsed[0]["data"]["cgpa"] == 8.5


def test_the_export_freezes_the_header_row():
    """A 400-row sheet is unusable without it."""
    from openpyxl import load_workbook

    buffer = build_workbook(COLUMNS, [], "Students")
    assert load_workbook(buffer).active.freeze_panes == "A2"
