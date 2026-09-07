"""Shared helpers for importing/exporting domain records as .xlsx workbooks.

A list of :class:`Column` specs drives three things consistently: the export
sheet, the blank import template, and parsing/validating an uploaded file.
"""
import io
from typing import Any, Optional

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

_TRUTHY = {"true", "yes", "y", "1"}
_FALSY = {"false", "no", "n", "0", ""}


class Column:
    """Describes one spreadsheet column and how it maps to a model attribute."""

    def __init__(
        self,
        header: str,
        attr: str,
        kind: str = "str",
        required: bool = False,
        importable: bool = True,
        choices: Optional[list[str]] = None,
    ):
        self.header = header
        self.attr = attr
        self.kind = kind  # str | int | float | bool | enum
        self.required = required
        self.importable = importable  # False => export/template only
        self.choices = choices  # allowed lowercased values for kind == "enum"


def _is_blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and value.strip() == "")


def _parse_value(raw: Any, col: Column) -> Any:
    if _is_blank(raw):
        return None
    if col.kind == "str":
        return str(raw).strip()
    if col.kind == "int":
        try:
            return int(float(raw))
        except (TypeError, ValueError):
            raise ValueError(f"expected a whole number, got '{raw}'")
    if col.kind == "float":
        try:
            return float(raw)
        except (TypeError, ValueError):
            raise ValueError(f"expected a number, got '{raw}'")
    if col.kind == "bool":
        if isinstance(raw, bool):
            return raw
        token = str(raw).strip().lower()
        if token in _TRUTHY:
            return True
        if token in _FALSY:
            return False
        raise ValueError(f"expected yes/no, got '{raw}'")
    if col.kind == "enum":
        token = str(raw).strip().lower()
        if col.choices and token not in col.choices:
            raise ValueError(f"'{raw}' is not one of {', '.join(col.choices)}")
        return token
    return raw


def build_workbook(columns: list[Column], rows: list[Any], sheet_title: str) -> io.BytesIO:
    """Render ``rows`` (model instances) into an in-memory .xlsx using ``columns``."""
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="4F46E5")
    for i, col in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=i, value=col.header)
        cell.font = header_font
        cell.fill = header_fill
        ws.column_dimensions[get_column_letter(i)].width = max(14, len(col.header) + 2)

    for r, obj in enumerate(rows, start=2):
        for c, col in enumerate(columns, start=1):
            value = getattr(obj, col.attr, None)
            if hasattr(value, "value"):  # unwrap enum members
                value = value.value
            ws.cell(row=r, column=c, value=value)

    ws.freeze_panes = "A2"
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def _match_headers(header_row: Any, importable: dict[str, Column]) -> dict[str, tuple[int, Column]]:
    """Map ``attr -> (column index, Column)`` for the known headers in a row."""
    if not header_row:
        return {}
    headers = [str(h).strip().lower() if h is not None else "" for h in header_row]
    return {col.attr: (i, col) for i, h in enumerate(headers) if (col := importable.get(h))}


def parse_rows(file_bytes: bytes, columns: list[Column]) -> list[dict]:
    """Parse an uploaded workbook into ``[{"row", "data", "errors"}, ...]``.

    The sheet is chosen by looking for the template's headers rather than by
    taking whichever tab happened to be active when the file was saved — that
    silently imported nothing when someone kept their data on a second tab.
    Headers are matched case-insensitively; unknown columns are ignored and
    fully blank rows are skipped. Per-cell validation problems are collected
    into ``errors`` rather than raised, so one bad row never aborts the batch.

    Raises ``ValueError`` if no sheet carries the expected header row, so the
    caller can tell the user what's wrong instead of reporting "0 imported".
    """
    try:
        wb = load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
    except Exception as exc:  # noqa: BLE001 - surface a friendly message to the API caller
        raise ValueError(f"Could not read file as .xlsx: {exc}") from exc

    importable = {c.header.lower(): c for c in columns if c.importable}
    active_title = wb.active.title if wb.active is not None else None

    # Score every sheet by how many known columns its first row names, then by
    # whether it actually holds any data — a filled second tab must beat the
    # empty template tab even when both carry the same headers. Remaining ties go
    # to the active sheet, then to the leftmost tab.
    best_ws = None
    best_key = ()
    best_indexed: dict[str, tuple[int, Column]] = {}
    for position, sheet in enumerate(wb.worksheets):
        rows = sheet.iter_rows(values_only=True)
        indexed = _match_headers(next(rows, None), importable)
        if not indexed:
            continue
        has_data = any(not all(_is_blank(v) for v in row) for row in rows)
        key = (len(indexed), has_data, sheet.title == active_title, -position)
        if key > best_key:
            best_ws, best_key, best_indexed = sheet, key, indexed

    if best_ws is None:
        required = [c.header for c in columns if c.importable and c.required]
        expected = required or [c.header for c in columns if c.importable][:3]
        sheets = ", ".join(f"'{s.title}'" for s in wb.worksheets) or "none"
        raise ValueError(
            "Could not find the import columns in this file. The first row of a "
            f"sheet must hold the template's column headers (e.g. {', '.join(expected)}). "
            f"Sheets checked: {sheets}."
        )

    # A required column missing from the header row is a problem with the file,
    # not with each row in it — report it once rather than repeating it per row.
    missing = [c.header for c in columns if c.importable and c.required and c.attr not in best_indexed]
    if missing:
        raise ValueError(
            f"Sheet '{best_ws.title}' is missing the required column(s): "
            f"{', '.join(missing)}. Header names must match the template exactly."
        )

    # Re-read the winning sheet from the top; read-only iteration is one-shot.
    all_rows = list(best_ws.iter_rows(values_only=True))
    indexed = best_indexed

    parsed: list[dict] = []
    for row_number, raw_row in enumerate(all_rows[1:], start=2):
        if all(_is_blank(v) for v in raw_row):
            continue
        data: dict[str, Any] = {}
        errors: list[str] = []
        for attr, (i, col) in indexed.items():
            raw = raw_row[i] if i < len(raw_row) else None
            try:
                value = _parse_value(raw, col)
            except ValueError as exc:
                errors.append(f"{col.header}: {exc}")
                continue
            if value is None and col.required:
                errors.append(f"{col.header} is required")
                continue
            data[attr] = value
        parsed.append({"row": row_number, "data": data, "errors": errors})
    return parsed
