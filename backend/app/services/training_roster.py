"""Read an uploaded training attendance sheet (.xlsx) into normalized rows.

The same problem as ``round_roster`` and solved the same way: colleges keep
attendance in whatever shape the trainer handed over, so map the columns by
header name and accept the aliases people actually type. Roll number is the only
required column — everything else is optional, because a bare list of who turned
up is still worth importing.

Deliberately no AI here. A roster sheet can be an arbitrary layout worth paying a
model to interpret; an attendance sheet is a column of roll numbers and some
marks, and if the roll-number column can't be found the honest answer is to say
so rather than spend credit guessing.
"""

import io
from typing import Optional

from openpyxl import load_workbook

from app.services.round_roster import _NAME_ALIASES, _ROLL_ALIASES

_ATTENDANCE_ALIASES = {
    "attendance", "attendance %", "attendance percent", "attendance percentage",
    "attendance_percent", "present %", "present", "attendance (%)", "att %", "att",
}
_SCORE_ALIASES = {
    "score", "marks", "result", "final score", "assessment", "assessment score",
    "total", "marks obtained", "score (%)", "grade",
}
_MOCK_ALIASES = {
    "mock", "mock score", "mock test", "mock test score", "mock_test_score",
    "mock interview", "mock interview score", "interview score",
}
_STATUS_ALIASES = {"status", "completion", "completed", "outcome", "remarks", "remark"}

# Free text in a status column, mapped to the four states a record can hold.
_COMPLETED = {"completed", "complete", "done", "pass", "passed", "finished", "yes", "y", "cleared"}
_DROPPED = {"dropped", "drop", "withdrawn", "withdrew", "left", "discontinued", "absent", "no show"}
_IN_PROGRESS = {"in progress", "in_progress", "ongoing", "attending", "started", "partial"}

TEMPLATE_HEADERS = ["Roll Number", "Name", "Attendance %", "Score", "Mock Score", "Status"]


def normalize_status(raw) -> Optional[str]:
    if raw is None or (isinstance(raw, str) and raw.strip() == ""):
        return None
    token = str(raw).strip().lower()
    if token in _COMPLETED:
        return "completed"
    if token in _DROPPED:
        return "dropped"
    if token in _IN_PROGRESS:
        return "in_progress"
    return None


def _to_float(raw) -> Optional[float]:
    """Numbers arrive as "85", "85%", "85 %" or a real number."""
    if raw is None or (isinstance(raw, str) and raw.strip() == ""):
        return None
    if isinstance(raw, (int, float)):
        return float(raw)
    token = str(raw).strip().rstrip("%").strip()
    try:
        return float(token)
    except (TypeError, ValueError):
        return None


def _find(headers: list[str], aliases: set[str]) -> Optional[int]:
    for i, h in enumerate(headers):
        if h in aliases:
            return i
    return None


def extract_attendance(file_bytes: bytes) -> list[dict]:
    """Rows of ``{roll_number, name, attendance_percent, score, mock_test_score,
    status}``. Raises ValueError with a message worth showing the user when the
    file can't be read or has no recognisable roll-number column."""
    try:
        wb = load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
    except Exception as exc:  # noqa: BLE001 - friendly message for the API caller
        raise ValueError(f"Could not read that file as .xlsx: {exc}") from exc

    ws = wb.active
    all_rows = [list(r) for r in ws.iter_rows(values_only=True)]
    if not all_rows:
        raise ValueError("That sheet is empty.")

    headers = [str(h).strip().lower() if h is not None else "" for h in all_rows[0]]
    roll_i = _find(headers, _ROLL_ALIASES)
    if roll_i is None:
        raise ValueError(
            "No roll number column found. Name one column 'Roll Number' (or "
            "'Reg No', 'USN') and upload again — the template has the right headers."
        )

    name_i = _find(headers, _NAME_ALIASES)
    attendance_i = _find(headers, _ATTENDANCE_ALIASES)
    score_i = _find(headers, _SCORE_ALIASES)
    mock_i = _find(headers, _MOCK_ALIASES)
    status_i = _find(headers, _STATUS_ALIASES)

    def cell(row: list, i: Optional[int]):
        return row[i] if i is not None and i < len(row) else None

    rows: list[dict] = []
    seen: set[str] = set()
    for raw in all_rows[1:]:
        roll = cell(raw, roll_i)
        if roll is None or str(roll).strip() == "":
            continue
        roll_number = str(roll).strip()
        # A roll number twice in one sheet is a duplicate line, not two students;
        # the first wins so a re-paste can't overwrite good marks with blanks.
        key = roll_number.lower()
        if key in seen:
            continue
        seen.add(key)
        name = cell(raw, name_i)
        rows.append({
            "roll_number": roll_number,
            "name": str(name).strip() if name not in (None, "") else None,
            "attendance_percent": _to_float(cell(raw, attendance_i)),
            "score": _to_float(cell(raw, score_i)),
            "mock_test_score": _to_float(cell(raw, mock_i)),
            "status": normalize_status(cell(raw, status_i)),
        })
    return rows
