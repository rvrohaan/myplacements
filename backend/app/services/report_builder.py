"""Rendering for the reports in ``services.report_defs``.

A report is described once — title, the parameters it was run with, a list of
sections, and the caveats that belong with the figures — and rendered two ways
from that single description: as JSON for the screen, and as an .xlsx for the
file somebody attaches to an email or files with an accreditation submission.

Doing it this way is the point: a report whose downloaded copy can disagree with
what was on screen is worse than no report, because the disagreement surfaces in
a meeting rather than here.
"""

import io
from dataclasses import dataclass, field
from datetime import datetime
from typing import Any, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

_TITLE = Font(bold=True, size=14)
_META = Font(size=9, color="666666")
_HEADING = Font(bold=True, size=11)
_HEADER = Font(bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill("solid", fgColor="2563EB")
_TOTAL = Font(bold=True)
_NOTE = Font(size=9, italic=True, color="666666")
_THIN = Side(style="thin", color="E5E7EB")
_CELL_BORDER = Border(bottom=_THIN)


@dataclass
class Section:
    """One table in a report."""

    heading: str
    columns: list[str]
    rows: list[list[Any]]
    #: What these figures do *not* say — the basis they count on, or why a
    #: correlation isn't a cause. Travels with the table into the workbook, so
    #: the caveat can't be lost by exporting it.
    note: Optional[str] = None
    #: Rendered bold and kept out of any sorting; the last row of a table that
    #: has one.
    total_row: Optional[list[Any]] = None


@dataclass
class Report:
    id: str
    title: str
    subtitle: Optional[str] = None
    #: Label/value pairs describing the run: college, batch, filters, when.
    meta: list[tuple[str, str]] = field(default_factory=list)
    sections: list[Section] = field(default_factory=list)
    #: Prose above the tables, when a report has any: paragraphs, not a table.
    #: It lives on the Report rather than inside a one-column Section so both
    #: renderings lay it out as prose — a summary that arrives in the workbook
    #: as a column of cells is not the same document that was on screen.
    narrative: list[str] = field(default_factory=list)
    #: Where the narrative came from, shown beside it. None when there is none.
    narrative_note: Optional[str] = None
    #: Limits of the report as a whole. Always rendered, never optional — a
    #: figure quoted without its basis is how a placement report becomes wrong.
    caveats: list[str] = field(default_factory=list)
    generated_at: datetime = field(default_factory=datetime.utcnow)

    def to_dict(self) -> dict:
        return {
            "id": self.id,
            "title": self.title,
            "subtitle": self.subtitle,
            "generated_at": self.generated_at.isoformat(),
            "meta": [{"label": k, "value": v} for k, v in self.meta],
            "sections": [
                {
                    "heading": s.heading,
                    "columns": s.columns,
                    "rows": [[_plain(v) for v in row] for row in s.rows],
                    "total_row": [_plain(v) for v in s.total_row] if s.total_row else None,
                    "note": s.note,
                }
                for s in self.sections
            ],
            "narrative": self.narrative,
            "narrative_note": self.narrative_note,
            "caveats": self.caveats,
        }


def _plain(value: Any) -> Any:
    """JSON-safe cell value. None stays None rather than becoming "None" or 0 —
    "not recorded" and "zero" are different answers and must stay different."""
    if isinstance(value, datetime):
        return value.isoformat()
    return value


def build_report_workbook(report: Report) -> io.BytesIO:
    """Render a report as a single-sheet .xlsx, laid out to be read and printed
    rather than parsed: a title block, then each section with its note."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Report"
    widths: dict[int, int] = {}
    row_at = 1

    def write(col: int, row: int, value: Any, font=None, wrap: bool = False, border=False):
        cell = ws.cell(row=row, column=col, value=value)
        if font:
            cell.font = font
        if wrap:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        if border:
            cell.border = _CELL_BORDER
        text = "" if value is None else str(value)
        # Wrapped cells shouldn't drag a column out to the width of a sentence.
        measured = min(len(text), 60) if not wrap else 0
        widths[col] = max(widths.get(col, 10), measured + 2)
        return cell

    write(1, row_at, report.title, _TITLE)
    row_at += 1
    if report.subtitle:
        write(1, row_at, report.subtitle, _META)
        row_at += 1
    row_at += 1

    for label, value in report.meta:
        write(1, row_at, label, _META)
        write(2, row_at, value, _META)
        row_at += 1
    row_at += 1

    # Prose first, the way it reads on screen. Merged across the table width so
    # a paragraph wraps as a paragraph instead of running off into column B.
    if report.narrative:
        write(1, row_at, "The month in words", _HEADING)
        row_at += 1
        for paragraph in report.narrative:
            ws.merge_cells(start_row=row_at, start_column=1, end_row=row_at, end_column=6)
            cell = write(1, row_at, paragraph, _NOTE, wrap=True)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            # Roughly one line per 110 characters across six merged columns.
            ws.row_dimensions[row_at].height = 15 * max(1, len(paragraph) // 110 + 1)
            row_at += 1
        if report.narrative_note:
            write(1, row_at, report.narrative_note, _NOTE)
            row_at += 1
        row_at += 1

    for section in report.sections:
        write(1, row_at, section.heading, _HEADING)
        row_at += 1
        if section.note:
            write(1, row_at, section.note, _NOTE)
            row_at += 1

        for i, header in enumerate(section.columns, start=1):
            cell = write(i, row_at, header, _HEADER)
            cell.fill = _HEADER_FILL
        row_at += 1

        for data_row in section.rows:
            for i, value in enumerate(data_row, start=1):
                write(i, row_at, _cell(value), border=True)
            row_at += 1

        if section.total_row:
            for i, value in enumerate(section.total_row, start=1):
                write(i, row_at, _cell(value), font=_TOTAL, border=True)
            row_at += 1
        row_at += 1

    if report.caveats:
        write(1, row_at, "How to read this", _HEADING)
        row_at += 1
        for caveat in report.caveats:
            write(1, row_at, f"• {caveat}", _NOTE, wrap=True)
            row_at += 1

    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = min(width, 62)
    ws.freeze_panes = "A1"

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def _cell(value: Any) -> Any:
    """A blank cell for "not recorded", so a spreadsheet total never silently
    counts it as zero."""
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.replace(tzinfo=None)
    return value
