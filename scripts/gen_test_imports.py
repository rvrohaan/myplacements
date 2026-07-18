"""Generate sample import workbooks for manual frontend testing.

Headers below match the Column specs in app/routers/companies.py and
app/routers/students.py exactly (matched case-insensitively on import).
"""
import os

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

OUT_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "test_data")
os.makedirs(OUT_DIR, exist_ok=True)


def write(filename, headers, rows, title):
    wb = Workbook()
    ws = wb.active
    ws.title = title
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="4F46E5")
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=1, column=i, value=h)
        c.font = header_font
        c.fill = header_fill
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = max(14, len(h) + 2)
    for r, row in enumerate(rows, start=2):
        for col, value in enumerate(row, start=1):
            ws.cell(row=r, column=col, value=value)
    ws.freeze_panes = "A2"
    path = os.path.join(OUT_DIR, filename)
    wb.save(path)
    print("wrote", path)


# --- Companies (10) ---
company_headers = [
    "name", "sector", "domain", "location", "size", "website",
    "products_services", "status", "mou_status", "hiring_pattern",
    "preferred_branches", "min_cgpa", "salary_min", "salary_max", "notes",
]
company_rows = [
    ["Infosys", "IT Services", "Consulting", "Bengaluru", "Large", "https://infosys.com",
     "Digital transformation, cloud, AI services", "active", "Signed", "Annual bulk hiring",
     "CSE, IT, ECE", 6.5, 4.5, 9.0, "Tier-1 mass recruiter"],
    ["TCS", "IT Services", "Consulting", "Mumbai", "Large", "https://tcs.com",
     "IT consulting and business solutions", "priority", "Signed", "NQT-based hiring",
     "All branches", 6.0, 3.6, 7.0, "Hires via National Qualifier Test"],
    ["Zoho", "Product", "SaaS", "Chennai", "Medium", "https://zoho.com",
     "Cloud-based business software suite", "active", "In progress", "On-campus drive",
     "CSE, IT", 7.0, 6.0, 10.0, "Strong product engineering culture"],
    ["Freshworks", "Product", "SaaS", "Chennai", "Medium", "https://freshworks.com",
     "Customer engagement software", "new", "Not started", "Selective hiring",
     "CSE, IT, ECE", 7.5, 8.0, 14.0, "Recently expanded campus program"],
    ["Wipro", "IT Services", "Consulting", "Bengaluru", "Large", "https://wipro.com",
     "IT, consulting and BPO services", "active", "Signed", "Elite & Turbo tracks",
     "CSE, IT, ECE, EEE", 6.0, 3.5, 8.0, "Multiple hiring tracks"],
    ["Amazon", "E-commerce", "Cloud/Retail", "Hyderabad", "Large", "https://amazon.jobs",
     "E-commerce and AWS cloud services", "priority", "Signed", "SDE hiring drive",
     "CSE, IT", 7.0, 18.0, 30.0, "Dream company, high CTC"],
    ["Razorpay", "Fintech", "Payments", "Bengaluru", "Medium", "https://razorpay.com",
     "Online payment gateway and banking", "new", "Not started", "Off-campus referrals",
     "CSE, IT", 7.5, 12.0, 20.0, "Fast-growing fintech"],
    ["Deloitte", "Consulting", "Audit/Advisory", "Gurugram", "Large", "https://deloitte.com",
     "Audit, tax, and consulting services", "active", "In progress", "Analyst roles",
     "All branches", 6.5, 5.0, 9.0, "Hires across functions"],
    ["Mindtree", "IT Services", "Digital", "Bengaluru", "Medium", "https://ltimindtree.com",
     "Digital and IT consulting", "dormant", "Expired", "Paused this cycle",
     "CSE, IT, ECE", 6.0, 4.0, 7.5, "No drive last two seasons"],
    ["Acme Corp", "Manufacturing", "Industrial", "Pune", "Small", "https://acme.example.com",
     "Industrial equipment manufacturing", "blacklisted", "Terminated", "None",
     "MECH, EEE", 5.5, 3.0, 5.0, "Payment defaults reported"],
]
write("companies_test_import.xlsx", company_headers, company_rows, "Companies")


# --- Students (20) ---
student_headers = [
    "full_name", "roll_number", "branch", "batch_year", "cgpa", "backlogs",
    "skills", "certifications", "internships", "projects", "resume_url",
    "linkedin_url", "github_url", "placement_preference", "location_preference",
    "higher_studies_plan",
]
first = ["Aarav", "Diya", "Vivaan", "Ananya", "Arjun", "Saanvi", "Reyansh", "Ishaan",
         "Myra", "Aditya", "Kiara", "Vihaan", "Aadhya", "Krish", "Anika", "Rohan",
         "Navya", "Karthik", "Pooja", "Siddharth"]
last = ["Sharma", "Patel", "Reddy", "Nair", "Iyer", "Gupta", "Mehta", "Singh",
        "Rao", "Kulkarni", "Das", "Bose", "Menon", "Joshi", "Verma", "Pillai",
        "Chopra", "Kumar", "Shetty", "Banerjee"]
branches = ["CSE", "IT", "ECE", "EEE", "MECH"]
skills_pool = ["Python, SQL, React", "Java, Spring Boot, MySQL", "C++, Data Structures",
               "JavaScript, Node.js, MongoDB", "Embedded C, VLSI", "Python, ML, Pandas",
               "Go, Kubernetes, AWS", "Power systems, MATLAB", "AutoCAD, SolidWorks",
               "Flutter, Dart, Firebase"]
prefs = ["Software", "Data Science", "Core", "Consulting", "Product Management"]
locations = ["Bengaluru", "Hyderabad", "Chennai", "Pune", "Remote", "Mumbai"]

student_rows = []
for i in range(20):
    branch = branches[i % len(branches)]
    cgpa = round(6.0 + (i % 8) * 0.45, 2)
    backlogs = i % 3
    higher = "yes" if i % 5 == 0 else "no"
    roll = f"21{branch[:2].upper()}{1001 + i}"
    name = f"{first[i]} {last[i]}"
    student_rows.append([
        name, roll, branch, 2025 + (i % 2), cgpa, backlogs,
        skills_pool[i % len(skills_pool)],
        "AWS Cloud Practitioner" if i % 2 == 0 else "",
        f"Summer intern at TestCo ({2 + i % 3} months)" if i % 3 != 0 else "",
        f"Built a {prefs[i % len(prefs)].lower()} capstone project",
        f"https://resumes.example.com/{roll}.pdf",
        f"https://linkedin.com/in/{first[i].lower()}-{last[i].lower()}",
        f"https://github.com/{first[i].lower()}{last[i].lower()}",
        prefs[i % len(prefs)],
        locations[i % len(locations)],
        higher,
    ])
write("students_test_import.xlsx", student_headers, student_rows, "Students")
