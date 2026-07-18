"""Generate an .xlsx with 10 more fake companies, ready to import.

Headers match the Column specs in app/routers/companies.py exactly (matched
case-insensitively on import). `status` must be one of the CompanyStatus enum
values: active, dormant, blacklisted, priority, new.

Run:  python scripts/gen_more_companies.py
Output: test_data/companies_more_import.xlsx
"""
import os

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

OUT_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "test_data")
os.makedirs(OUT_DIR, exist_ok=True)

headers = [
    "name", "sector", "domain", "location", "size", "website",
    "products_services", "status", "mou_status", "hiring_pattern",
    "preferred_branches", "min_cgpa", "salary_min", "salary_max", "notes",
]

rows = [
    ["HCLTech", "IT Services", "Engineering", "Noida", "Large", "https://hcltech.com",
     "IT and engineering R&D services", "active", "Signed", "Annual campus drive",
     "CSE, IT, ECE, MECH", 6.0, 3.8, 7.5, "Consistent bulk recruiter"],
    ["Flipkart", "E-commerce", "Retail Tech", "Bengaluru", "Large", "https://flipkart.com",
     "Online marketplace and logistics tech", "priority", "In progress", "SDE + PM roles",
     "CSE, IT", 7.5, 16.0, 28.0, "Dream company, competitive process"],
    ["PhonePe", "Fintech", "Payments", "Bengaluru", "Medium", "https://phonepe.com",
     "UPI payments and financial services", "new", "Not started", "Off-campus hiring",
     "CSE, IT", 7.5, 14.0, 24.0, "High-growth fintech"],
    ["Persistent Systems", "IT Services", "Digital", "Pune", "Large", "https://persistent.com",
     "Software product engineering services", "active", "Signed", "On-campus drive",
     "CSE, IT, ECE", 6.5, 4.5, 8.5, "Steady annual hiring"],
    ["Swiggy", "Food Tech", "Delivery", "Bengaluru", "Medium", "https://swiggy.com",
     "Food and grocery delivery platform", "new", "Not started", "Selective SDE hiring",
     "CSE, IT", 7.0, 12.0, 22.0, "Product-focused engineering"],
    ["Cognizant", "IT Services", "Consulting", "Chennai", "Large", "https://cognizant.com",
     "IT consulting and digital services", "active", "Signed", "GenC hiring program",
     "All branches", 6.0, 4.0, 6.5, "Large-volume entry-level intake"],
    ["Qualcomm", "Semiconductor", "Chip Design", "Hyderabad", "Large", "https://qualcomm.com",
     "Wireless chipsets and VLSI design", "priority", "In progress", "Core VLSI roles",
     "ECE, EEE", 7.5, 14.0, 26.0, "Top core-electronics recruiter"],
    ["Nagarro", "IT Services", "Digital", "Gurugram", "Medium", "https://nagarro.com",
     "Digital product engineering", "active", "Signed", "Fresher hiring drive",
     "CSE, IT, ECE", 6.5, 5.0, 9.0, "Good mid-tier package"],
    ["BharatBenz", "Manufacturing", "Automotive", "Chennai", "Large", "https://bharatbenz.com",
     "Commercial vehicle manufacturing", "dormant", "Expired", "Paused hiring",
     "MECH, EEE, ECE", 6.0, 4.0, 6.5, "No drive last cycle"],
    ["Zeta Solutions", "Startup", "SaaS", "Remote", "Small", "https://zeta.example.com",
     "Early-stage B2B SaaS tooling", "new", "Not started", "Intern-to-hire",
     "CSE, IT", 7.0, 6.0, 12.0, "Seed-stage, equity component"],
]

wb = Workbook()
ws = wb.active
ws.title = "Companies"
header_font = Font(bold=True, color="FFFFFF")
header_fill = PatternFill("solid", fgColor="4F46E5")
for i, h in enumerate(headers, start=1):
    c = ws.cell(row=1, column=i, value=h)
    c.font = header_font
    c.fill = header_fill
    ws.column_dimensions[c.column_letter].width = max(14, len(h) + 2)
for r, row in enumerate(rows, start=2):
    for col, value in enumerate(row, start=1):
        ws.cell(row=r, column=col, value=value)
ws.freeze_panes = "A2"

path = os.path.join(OUT_DIR, "companies_more_import.xlsx")
wb.save(path)
print("wrote", path)
